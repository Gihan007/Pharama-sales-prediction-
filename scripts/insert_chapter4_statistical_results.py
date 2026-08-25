from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOC = ROOT / "docs" / "thesisi version" / "CSCI 43018- final unstructiured draft (1) - current state 2_new.docx"
METRICS = ROOT / "src" / "evaluation_results" / "model_metrics.json"
TIMES = ROOT / "src" / "evaluation_results" / "inference_times.json"
IMAGE_DIR = ROOT / "src" / "static" / "images"
OUT_DIR = ROOT / "docs" / "thesisi version" / "chapter4_inserted_assets"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    node = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "6")
        item.set(qn("w:color"), "808080")
        node.append(item)
    tbl_pr.append(node)


def format_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    borders(table)
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r == 0:
                shade(cell, "1F4E78")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9)
                    if r == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def caption(document, text: str):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    return p


def picture_block(document, image: Path, caption_text: str):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image), width=Inches(6.25))
    return [p._p, caption(document, caption_text)._p]


def table_block(document, headers, rows, caption_text: str):
    cap = caption(document, caption_text)
    table = document.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    format_table(table)
    return [cap._p, table._tbl]


def insert_after(paragraph, elements) -> None:
    cursor = paragraph._p
    for element in elements:
        cursor.addnext(element)
        cursor = element


def find_para(document, startswith: str):
    for p in document.paragraphs:
        if " ".join(p.text.split()).startswith(startswith):
            return p
    raise RuntimeError(f"Anchor paragraph not found: {startswith}")


def make_mae_chart(metrics: dict) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / "figure_4_3_category_mae_comparison.png"
    cats = list(metrics)
    sarimax = [metrics[c]["sarimax"]["MAE"] for c in cats]
    prophet = [metrics[c]["prophet"]["MAE"] for c in cats]
    x = np.arange(len(cats))
    fig, ax = plt.subplots(figsize=(9.4, 4.8), dpi=180)
    width = 0.38
    ax.bar(x - width / 2, sarimax, width, label="SARIMAX", color="#1f77b4")
    ax.bar(x + width / 2, prophet, width, label="Prophet", color="#e377c2")
    ax.set_ylabel("Mean Absolute Error (MAE)")
    ax.set_xlabel("Drug category")
    ax.set_xticks(x, cats)
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def main() -> None:
    metrics = json.loads(METRICS.read_text(encoding="utf-8"))
    times = json.loads(TIMES.read_text(encoding="utf-8"))
    document = Document(DOC)
    marker = "Table 4.1: Category-Wise Performance of SARIMAX and Prophet"
    if any(marker in p.text for p in document.paragraphs):
        raise RuntimeError("Chapter 4 statistical results have already been inserted")

    cats = list(metrics)
    category_rows = []
    for c in cats:
        s, p = metrics[c]["sarimax"], metrics[c]["prophet"]
        category_rows.append([
            c, f'{s["MAE"]:.4f}', f'{s["RMSE"]:.4f}', f'{s["MAPE"]:.4f}',
            f'{p["MAE"]:.4f}', f'{p["RMSE"]:.4f}', f'{p["MAPE"]:.4f}',
        ])

    avg_rows = []
    for model, name in (("sarimax", "SARIMAX"), ("prophet", "Prophet")):
        avg_rows.append([
            name,
            f'{np.mean([metrics[c][model]["MAE"] for c in cats]):.4f}',
            f'{np.mean([metrics[c][model]["RMSE"] for c in cats]):.4f}',
            f'{np.mean([metrics[c][model]["MAPE"] for c in cats]):.4f}',
            f'{np.mean([times[c][model] for c in cats]):.4f}',
        ])

    setup_anchor = find_para(document, "The statistical models were evaluated for all eight drug categories.")
    sarimax_anchor = find_para(document, "The results show that SARIMAX is useful as a baseline")
    prophet_anchor = find_para(document, "Across all categories, Prophet achieved an average MAE")
    summary_anchor = find_para(document, "Prophet 32.5682 37.6763 88.0285 0.8725 s")

    category_elements = table_block(
        document,
        ["Category", "SARIMAX MAE", "SARIMAX RMSE", "SARIMAX MAPE (%)", "Prophet MAE", "Prophet RMSE", "Prophet MAPE (%)"],
        category_rows,
        marker,
    )
    insert_after(setup_anchor, category_elements)

    insert_after(sarimax_anchor, picture_block(document, IMAGE_DIR / "C1_sarimax_forecast.png", "Figure 4.1: C1 Drug Sales Forecast Using the SARIMAX Model"))
    insert_after(prophet_anchor, picture_block(document, IMAGE_DIR / "C1_prophet_forecast.png", "Figure 4.2: C1 Drug Sales Forecast Using the Prophet Model"))

    avg_elements = table_block(
        document,
        ["Model", "Average MAE", "Average RMSE", "Average MAPE (%)", "Average Inference Time (s)"],
        avg_rows,
        "Table 4.2: Average Performance Comparison of Statistical Forecasting Models",
    )
    chart = make_mae_chart(metrics)
    insert_after(summary_anchor, avg_elements + picture_block(document, chart, "Figure 4.3: Category-Wise MAE Comparison of SARIMAX and Prophet"))

    document.save(DOC)
    print(DOC)


if __name__ == "__main__":
    main()
