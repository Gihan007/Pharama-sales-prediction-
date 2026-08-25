from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOC = ROOT / "docs" / "thesisi version" / "CSCI 43018- final unstructiured draft (1) - current state 2_new_RENUMBERED.docx"


EARLY_CAPTIONS = {
    "Figure 1 : Drug Category Classification Used in the Study": "Figure 1.1: Drug Category Classification Used in the Study",
    "FIGURE 2 : Data Preprocessing PipelinE": "Figure 3.1: Data Preprocessing Pipeline",
    "FIGURE 3 : Drug Category Classification Used in the Study": "Figure 3.2: Drug Category Classification Used in the Study",
    "Figure 4: Time-Series Feature Engineering Process": "Figure 3.3: Time-Series Feature Engineering Process",
    "Figure 5 : Time-Series Train-Test Splitting Method": "Figure 3.4: Time-Series Train-Test Splitting Method",
    "Figure 6: Forecasting Model Development Workflow": "Figure 3.5: Forecasting Model Development Workflow",
    "Figure 7 : Model Training and Artifact Storage Flow": "Figure 3.6: Model Training and Artifact Storage Flow",
    "Figure 8 : Forecasting Request-Response Flow": "Figure 3.7: Forecasting Request-Response Flow",
    "Figure 9: Explainability Workflow for Forecasting Outputs": "Figure 3.8: Explainability Workflow for Forecasting Outputs",
    "Figure 10: Advanced AI Module Workflow": "Figure 3.9: Advanced AI Module Workflow",
    "Figure 11: Microservice-Based System Architecture": "Figure 3.10: Microservice-Based System Architecture",
    "Figure 12 : API Gateway Communication Flow": "Figure 3.11: API Gateway Communication Flow",
}


def norm(text: str) -> str:
    return " ".join(text.split())


def format_caption(p) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = False
    for run in p.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.bold = True


def format_list_heading(p) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        run.bold = True


def list_paragraph(text: str):
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9360")
    tabs.append(tab); ppr.append(tabs)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    ppr.append(spacing); p.append(ppr)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman"); fonts.set(qn("w:hAnsi"), "Times New Roman")
    size = OxmlElement("w:sz"); size.set(qn("w:val"), "24")
    rpr.append(fonts); rpr.append(size); r.append(rpr)
    t = OxmlElement("w:t"); t.text = text; r.append(t); p.append(r)
    return p


def replace_between(start_p, end_p, entries: list[str]) -> None:
    parent = start_p._p.getparent()
    children = list(parent)
    start_i, end_i = children.index(start_p._p), children.index(end_p._p)
    for element in children[start_i + 1:end_i]:
        parent.remove(element)
    cursor = start_p._p
    for entry in entries:
        node = list_paragraph(entry + "\t")
        cursor.addnext(node)
        cursor = node


def main() -> None:
    document = Document(DOC)
    intro_i = next(i for i, p in enumerate(document.paragraphs) if norm(p.text) == "Introduction")

    # Normalize actual captions only, excluding obsolete front-matter list entries.
    for p in document.paragraphs[intro_i + 1:]:
        text = norm(p.text)
        if text in EARLY_CAPTIONS:
            p.text = EARLY_CAPTIONS[text]
            text = p.text
        if re.match(r"^(Figure|Table)\s+\d+\.\d+:\s+", text):
            p.text = text
            format_caption(p)

    actual = [norm(p.text) for p in document.paragraphs[intro_i + 1:]]
    figures = [t for t in actual if re.match(r"^Figure\s+\d+\.\d+:\s+", t)]
    tables = [t for t in actual if re.match(r"^Table\s+\d+\.\d+:\s+", t)]
    # The document has 37 inline images; one is the title-page logo and is not a numbered figure.
    if len(figures) != 36 or len(tables) != 17:
        raise RuntimeError(f"Unexpected caption count: {len(figures)} figures, {len(tables)} tables")

    list_tables = next(p for p in document.paragraphs if norm(p.text) == "List of Tables")
    list_figures = next(p for p in document.paragraphs if norm(p.text) == "List of Figures")
    acronyms = next(p for p in document.paragraphs if norm(p.text) == "List of Acronyms/Abbreviations")
    format_list_heading(list_tables); format_list_heading(list_figures)
    replace_between(list_tables, list_figures, tables)
    replace_between(list_figures, acronyms, figures)
    document.save(DOC)
    print(f"figures={len(figures)} tables={len(tables)}")


if __name__ == "__main__":
    main()
