from __future__ import annotations

import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document

from insert_chapter4_statistical_results import find_para, insert_after, picture_block, table_block


ROOT = Path(__file__).resolve().parents[1]
DOC = ROOT / "docs" / "thesisi version" / "CSCI 43018- final unstructiured draft (1) - current state 2_new.docx"
TEMP = ROOT / "docs" / "thesisi version" / "CSCI 43018- final unstructiured draft (1) - current state 2_new_with_4.10.docx"
SUMMARY = ROOT / "src" / "evaluation_results" / "performance_summary.json"
ENSEMBLE = ROOT / "src" / "evaluation_results" / "ensemble_results.json"
ASSET_DIR = ROOT / "docs" / "thesisi version" / "chapter4_inserted_assets"


def comparison_chart(rows: list[list[str]]) -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / "figure_4_23_overall_mae_rmse_comparison.png"
    labels = [r[1] for r in rows]
    mae = [float(r[2]) for r in rows]
    rmse = [float(r[3]) for r in rows]
    x = np.arange(len(labels)); width = 0.38
    fig, ax = plt.subplots(figsize=(10.2, 5.2), dpi=180)
    ax.bar(x - width / 2, mae, width, label="Average MAE", color="#2563eb")
    ax.bar(x + width / 2, rmse, width, label="Average RMSE", color="#d97706")
    ax.set_ylabel("Forecasting error")
    ax.set_xticks(x, labels, rotation=28, ha="right")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def main() -> None:
    document = Document(DOC)

    # Convert the existing plain-text 48/48 summary into a real Word table.
    plain = [
        "Testing Area Result",
        "Total endpoints tested 48",
        "Passed endpoints 48",
        "Failed endpoints 0",
        "Final status Successful",
    ]
    if any(" ".join(p.text.split()) == plain[0] for p in document.paragraphs):
        anchor = next(p for p in document.paragraphs if " ".join(p.text.split()) == plain[0])
        elements = table_block(
            document,
            ["Testing Area", "Result"],
            [["Total endpoints tested", "48"], ["Passed endpoints", "48"], ["Failed endpoints", "0"], ["Final status", "Successful"]],
            "Table 4.15: Reported Final API Testing Summary",
        )
        insert_after(anchor, elements)
        for p in list(document.paragraphs):
            if " ".join(p.text.split()) in set(plain):
                p._element.getparent().remove(p._element)

    marker = "Table 4.16: Overall Average Performance of Forecasting Models"
    if any(marker in p.text for p in document.paragraphs):
        raise RuntimeError("Section 4.10 results have already been inserted")

    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    ensemble = json.loads(ENSEMBLE.read_text(encoding="utf-8"))
    rows = []
    model_info = [
        ("Statistical", "SARIMAX", "sarimax"), ("Statistical", "Prophet", "prophet"),
        ("Machine learning", "XGBoost", "xgboost"), ("Machine learning", "LightGBM", "lightgbm"),
        ("Deep learning", "LSTM", "lstm"), ("Deep learning", "GRU", "gru"),
        ("Deep learning", "Transformer", "transformer"),
    ]
    for approach, label, key in model_info:
        item = summary[key]
        rows.append([approach, label, f'{item["final_mae"]:.4f}', f'{item["final_rmse"]:.4f}', f'{item["final_mape"]:.4f}', f'{item["final_time"]:.4f}'])
    for key, label in (("weighted_average", "Weighted Average"), ("performance_weighted", "Performance Weighted")):
        rows.append(["Ensemble", label,
            f'{np.mean([ensemble[c][key]["MAE"] for c in ensemble]):.4f}',
            f'{np.mean([ensemble[c][key]["RMSE"] for c in ensemble]):.4f}',
            f'{np.mean([ensemble[c][key]["MAPE"] for c in ensemble]):.4f}',
            f'{np.mean([ensemble[c][key]["Time"] for c in ensemble]):.4f}'])

    intro_anchor = find_para(document, "This section discusses the best-performing forecasting models based on the experimental results.")
    overall_anchor = find_para(document, "Considering all models, LightGBM was the best overall forecasting approach")
    insert_after(intro_anchor, table_block(
        document,
        ["Approach", "Model", "Average MAE", "Average RMSE", "Average MAPE (%)", "Average Time (s)"],
        rows,
        marker,
    ))
    insert_after(overall_anchor, picture_block(document, comparison_chart(rows), "Figure 4.23: Overall Average MAE and RMSE Comparison of Forecasting Models"))
    document.save(TEMP)
    print(TEMP)


if __name__ == "__main__":
    main()
